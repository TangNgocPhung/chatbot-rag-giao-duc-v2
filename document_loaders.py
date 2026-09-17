"""
BỘ LOADER THEO ĐỊNH DẠNG CHO KHO TÀI LIỆU GIÁO DỤC
====================================================
Mỗi định dạng có một loader riêng đã kiểm chứng trên kho thật, thay vì để
UnstructuredFileLoader đoán (thư viện này treo/crash với .docx và .pptx trên
máy Windows của dự án - xem ghi chú từng lớp bên dưới).

Loader trả về Document kèm metadata mô tả "vị trí" trong tài liệu gốc
(slide mấy, sheet nào, phút thứ mấy) để câu trả lời trích dẫn được chính xác:

    metadata["context_label"]  nhãn hiển thị cho người dùng, ví dụ "Slide 4"
    metadata["chunk_mode"]     "theo_phan_tu" nếu mỗi Document là một đơn vị
                               độc lập (slide/sheet/đoạn phiên âm) - bộ chunk
                               sẽ giữ nguyên ranh giới này thay vì gộp cả file
    metadata["time_start"]     giây bắt đầu (chỉ có ở video/âm thanh)
    metadata["so_trang"]       số trang PDF, đếm từ 1 (khớp neo #page= của
                               trình xem PDF, để "xem đoạn gốc" mở đúng trang)
"""

from __future__ import annotations

import os
import re
from html import unescape

from langchain_core.documents import Document

from media_transcribe import (
    DINH_DANG_MEDIA,
    dinh_dang_thoi_gian,
    doc_cache,
    doc_phu_de,
    phien_am,
)

DINH_DANG_VAN_BAN = {".pdf", ".docx", ".doc", ".txt", ".md", ".html", ".htm", ".epub"}
DINH_DANG_ANH = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
DINH_DANG_BANG = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv"}
DINH_DANG_TRINH_CHIEU = {".pptx"}
DINH_DANG_PHU_DE = {".srt", ".vtt"}
DINH_DANG_AM_THANH = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".wma"}

DINH_DANG_HO_TRO = sorted(
    DINH_DANG_VAN_BAN | DINH_DANG_BANG | DINH_DANG_TRINH_CHIEU
    | DINH_DANG_PHU_DE | DINH_DANG_MEDIA | DINH_DANG_ANH
)

# Số dòng dữ liệu bảng gộp vào một Document (header luôn được lặp lại).
SO_DONG_MOI_KHOI_BANG = 40
# Độ dài mục tiêu của một đoạn phiên âm gộp, tính theo ký tự.
DO_DAI_DOAN_PHIEN_AM = 900


class ThieuCongCuChuyenDoi(RuntimeError):
    """Không có Word/LibreOffice để đọc định dạng .doc đời cũ."""


def suy_loai_tai_lieu(duong_dan: str) -> str:
    """Nhóm định dạng dùng cho hiển thị và lọc trên giao diện."""
    duoi = os.path.splitext(duong_dan)[1].lower()
    if duoi in DINH_DANG_AM_THANH:
        return "am_thanh"
    if duoi in DINH_DANG_MEDIA:
        return "video"
    if duoi in DINH_DANG_BANG:
        return "bang_du_lieu"
    if duoi in DINH_DANG_TRINH_CHIEU:
        return "trinh_chieu"
    if duoi in DINH_DANG_PHU_DE:
        return "phu_de"
    if duoi in DINH_DANG_ANH:
        return "anh"
    return "van_ban"


# ============================================================
# PDF - ưu tiên lớp văn bản sẵn có, bản scan thì dùng OCR
# ============================================================
class PdfLoader:
    """
    Gần 80% PDF trong kho là bản scan ảnh (mỗi trang một tấm ảnh, 0 ký tự).
    PyPDFLoader trả về rỗng nên trước đây các văn bản quan trọng nhất - Luật
    Giáo dục, Nghị định, Thông tư - hoàn toàn không tra cứu được. Ở đây nếu
    phát hiện thiếu lớp văn bản thì lấy bản OCR (Tesseract, có cache).
    """

    def __init__(
        self, duong_dan: str, cho_phep_ocr: bool = True, bao_tien_do_ocr=None
    ):
        self.duong_dan = duong_dan
        self.cho_phep_ocr = cho_phep_ocr
        self.bao_tien_do_ocr = bao_tien_do_ocr

    def load(self):
        from langchain_community.document_loaders import PyPDFLoader

        import ocr_pdf

        tai_lieu = PyPDFLoader(self.duong_dan).load()
        du_van_ban = sum(len(d.page_content.strip()) for d in tai_lieu) >= (
            ocr_pdf.KY_TU_TOI_THIEU_MOI_TRANG * max(1, len(tai_lieu))
        )
        if du_van_ban:
            return self._danh_so_trang(tai_lieu)

        cac_trang = ocr_pdf.doc_cache(self.duong_dan)
        if cac_trang is None:
            if not self.cho_phep_ocr:
                return tai_lieu
            ok, thong_bao = ocr_pdf.san_sang()
            if not ok:
                # Không chặn cả lần cập nhật chỉ mục vì một file scan: trả về
                # phần văn bản ít ỏi đọc được, file sẽ bị đánh dấu no_text.
                return tai_lieu
            cac_trang = ocr_pdf.ocr_file_pdf(
                self.duong_dan, bao_tien_do=self.bao_tien_do_ocr
            )

        ket_qua = []
        for so_trang, van_ban in enumerate(cac_trang, 1):
            if not van_ban.strip():
                continue
            ket_qua.append(Document(
                page_content=van_ban,
                metadata={
                    "source": self.duong_dan,
                    "page": so_trang,
                    "so_trang": so_trang,
                    "nguon_van_ban": "ocr",
                },
            ))
        return ket_qua or tai_lieu

    @staticmethod
    def _danh_so_trang(tai_lieu):
        """PyPDFLoader đánh metadata["page"] từ 0, còn bản OCR ở trên đếm từ 1
        và neo #page= của trình xem PDF cũng đếm từ 1. Quy hết về một khóa
        "so_trang" đếm từ 1 để phía sau không phải nhớ chỗ nào lệch chỗ nào không."""
        for thu_tu, doc in enumerate(tai_lieu, 1):
            trang = doc.metadata.get("page")
            doc.metadata["so_trang"] = (
                trang + 1 if isinstance(trang, int) else thu_tu
            )
        return tai_lieu


# ============================================================
# TRÌNH CHIẾU (.pptx) - mỗi slide là một đơn vị độc lập
# ============================================================
class PptxLoader:
    """
    Trích text trực tiếp bằng python-pptx (không qua unstructured - toàn bộ
    52/52 file .pptx trong kho đều timeout/crash qua UnstructuredFileLoader).
    Mỗi slide thành một Document riêng kèm số slide + tiêu đề slide, nhờ vậy
    trích dẫn chỉ đúng slide thay vì cả bài giảng.
    """

    def __init__(self, duong_dan: str):
        self.duong_dan = duong_dan

    def load(self):
        import zipfile

        from pptx import Presentation

        # .pptx là một gói zip. Kho thật có file tải dở từ web (phần đuôi là
        # trang HTML báo lỗi), python-pptx chỉ báo "Package not found" rất khó
        # hiểu - nói thẳng ra là file hỏng để người dùng biết phải upload lại.
        if not zipfile.is_zipfile(self.duong_dan):
            raise ValueError(
                "File .pptx hỏng hoặc tải chưa đủ (không phải gói OpenXML hợp lệ). "
                "Hãy mở bằng PowerPoint và lưu lại, rồi upload lại lên Drive."
            )

        prs = Presentation(self.duong_dan)
        tai_lieu = []
        for so_slide, slide in enumerate(prs.slides, 1):
            cac_dong = []
            tieu_de = ""
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    for para in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in para.runs).strip()
                        if text:
                            cac_dong.append(text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        o = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if o:
                            cac_dong.append(" | ".join(o))
            try:
                if slide.shapes.title is not None and slide.shapes.title.text.strip():
                    tieu_de = slide.shapes.title.text.strip()
            except (AttributeError, ValueError):
                tieu_de = ""

            if slide.has_notes_slide:
                ghi_chu = (slide.notes_slide.notes_text_frame.text or "").strip()
                if ghi_chu:
                    cac_dong.append(f"Ghi chú của giáo viên: {ghi_chu}")

            if not cac_dong:
                continue
            nhan = f"Slide {so_slide}"
            if tieu_de and tieu_de != nhan:
                nhan = f"Slide {so_slide}: {tieu_de}"
            tai_lieu.append(Document(
                page_content=f"{nhan}\n" + "\n".join(cac_dong),
                metadata={
                    "source": self.duong_dan,
                    "context_label": nhan,
                    "slide": so_slide,
                    "chunk_mode": "theo_phan_tu",
                },
            ))
        return tai_lieu


# ============================================================
# BẢNG DỮ LIỆU (.xlsx/.xls/.csv) - mỗi sheet tách theo khối dòng
# ============================================================
class BangDuLieuLoader:
    """
    Đọc CSV/Excel theo từng sheet, cắt thành khối dòng và LẶP LẠI dòng tiêu đề
    ở mỗi khối. Nếu đổ cả workbook thành một khối văn bản rồi cắt theo ký tự
    (cách cũ), khối thứ hai trở đi mất header nên mô hình không còn biết cột
    nào là "Thứ", cột nào là "Tiết" - đọc sai thời khóa biểu.
    """

    def __init__(self, duong_dan: str):
        self.duong_dan = duong_dan

    def _doc_cac_bang(self):
        import pandas as pd

        duoi = os.path.splitext(self.duong_dan)[1].lower()
        if duoi in {".csv", ".tsv"}:
            ngan_cach = "\t" if duoi == ".tsv" else None
            bang = pd.read_csv(
                self.duong_dan, sep=ngan_cach, engine="python",
                dtype=str, keep_default_na=False, on_bad_lines="skip",
            )
            return {os.path.splitext(os.path.basename(self.duong_dan))[0]: bang}
        return pd.read_excel(self.duong_dan, sheet_name=None, dtype=str)

    @staticmethod
    def _bo_cot_va_dong_rong(bang):
        """File Excel hành chính thường thừa hàng chục cột/dòng trống."""
        van_ban = bang.astype(str).apply(lambda cot: cot.str.strip())
        bang = bang.loc[:, ~(van_ban == "").all(axis=0)]
        if bang.empty:
            return bang
        van_ban = bang.astype(str).apply(lambda cot: cot.str.strip())
        return bang[~(van_ban == "").all(axis=1)]

    def load(self):
        tai_lieu = []
        for ten_sheet, bang in self._doc_cac_bang().items():
            bang = self._bo_cot_va_dong_rong(bang.fillna(""))
            if bang.empty:
                continue

            dong_tieu_de = " | ".join(str(cot) for cot in bang.columns)
            tong_dong = len(bang)
            for bat_dau in range(0, tong_dong, SO_DONG_MOI_KHOI_BANG):
                khoi = bang.iloc[bat_dau: bat_dau + SO_DONG_MOI_KHOI_BANG]
                cac_dong = [
                    " | ".join(str(o).strip() for o in hang)
                    for hang in khoi.itertuples(index=False)
                ]
                ket_thuc = min(bat_dau + SO_DONG_MOI_KHOI_BANG, tong_dong)
                nhan = f"Sheet {ten_sheet}"
                if tong_dong > SO_DONG_MOI_KHOI_BANG:
                    nhan += f" · dòng {bat_dau + 1}-{ket_thuc}/{tong_dong}"
                tai_lieu.append(Document(
                    page_content=f"{nhan}\nCột: {dong_tieu_de}\n" + "\n".join(cac_dong),
                    metadata={
                        "source": self.duong_dan,
                        "context_label": nhan,
                        "sheet": str(ten_sheet),
                        "chunk_mode": "theo_phan_tu",
                    },
                ))
        return tai_lieu


# ============================================================
# HTML - dùng trafilatura (đã có sẵn cho tính năng đọc URL)
# ============================================================
class HtmlLoader:
    """Bóc phần nội dung chính, bỏ menu/quảng cáo; fallback sang gỡ thẻ thô."""

    def __init__(self, duong_dan: str):
        self.duong_dan = duong_dan

    def load(self):
        with open(self.duong_dan, "rb") as f:
            html = f.read().decode("utf-8", errors="replace")
        try:
            import trafilatura

            noi_dung = trafilatura.extract(
                html, include_comments=False, include_tables=True, no_fallback=False
            ) or ""
        except Exception:
            noi_dung = ""
        if len(noi_dung.strip()) < 40:
            khong_script = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
            noi_dung = re.sub(r"(?s)<[^>]+>", " ", khong_script)
            noi_dung = re.sub(r"\s+", " ", noi_dung).strip()
        return [Document(page_content=noi_dung, metadata={"source": self.duong_dan})]


class VanBanThuanLoader:
    def __init__(self, duong_dan: str):
        self.duong_dan = duong_dan

    def load(self):
        with open(self.duong_dan, encoding="utf-8", errors="replace") as f:
            return [Document(page_content=f.read(), metadata={"source": self.duong_dan})]


# ============================================================
# .doc ĐỜI CŨ - chuyển sang .docx rồi đọc
# ============================================================
def _chuyen_doc_sang_docx(duong_dan: str) -> str:
    """
    Word (COM) trước, LibreOffice sau. Trả về đường dẫn .docx tạm.
    docx2txt/python-docx KHÔNG đọc được .doc nhị phân đời cũ (OLE) nên bắt buộc
    phải chuyển đổi; 4 file .doc trong kho Drive đều thuộc dạng này.
    """
    import shutil
    import tempfile

    thu_muc_tam = tempfile.mkdtemp(prefix="rag-doc-")
    ten_goc = os.path.splitext(os.path.basename(duong_dan))[0]
    dich = os.path.join(thu_muc_tam, ten_goc + ".docx")

    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore

        # Cập nhật chỉ mục chạy trong luồng nền của web server, mà COM bắt buộc
        # phải khởi tạo riêng cho từng luồng - thiếu bước này Dispatch sẽ báo
        # "CoInitialize has not been called" và mọi file .doc đều thất bại.
        pythoncom.CoInitialize()
        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            try:
                tai_lieu = word.Documents.Open(os.path.abspath(duong_dan), ReadOnly=True)
                tai_lieu.SaveAs2(dich, FileFormat=16)  # 16 = wdFormatDocumentDefault
                tai_lieu.Close(False)
            finally:
                word.Quit()
        finally:
            pythoncom.CoUninitialize()
        if os.path.exists(dich):
            return dich
    except Exception:
        pass

    soffice = shutil.which("soffice") or shutil.which("soffice.exe")
    if not soffice:
        for ung_vien in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ):
            if os.path.exists(ung_vien):
                soffice = ung_vien
                break
    if soffice:
        import subprocess

        subprocess.run(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", thu_muc_tam,
             os.path.abspath(duong_dan)],
            capture_output=True, timeout=240, check=False,
        )
        if os.path.exists(dich):
            return dich

    shutil.rmtree(thu_muc_tam, ignore_errors=True)
    raise ThieuCongCuChuyenDoi(
        "Không đọc được .doc đời cũ: máy cần Microsoft Word (kèm pywin32) hoặc "
        "LibreOffice để chuyển sang .docx. Cách khác: mở file và lưu lại dạng .docx."
    )


class DocCuLoader:
    def __init__(self, duong_dan: str):
        self.duong_dan = duong_dan

    def load(self):
        import shutil

        from langchain_community.document_loaders import Docx2txtLoader

        duong_dan_docx = _chuyen_doc_sang_docx(self.duong_dan)
        try:
            tai_lieu = Docx2txtLoader(duong_dan_docx).load()
        finally:
            shutil.rmtree(os.path.dirname(duong_dan_docx), ignore_errors=True)
        for doc in tai_lieu:
            doc.metadata["source"] = self.duong_dan
            doc.metadata["da_chuyen_doi"] = "doc->docx"
        return tai_lieu


# ============================================================
# VIDEO / ÂM THANH / PHỤ ĐỀ - phiên âm thành văn bản có mốc thời gian
# ============================================================
def _gop_doan_phien_am(cac_doan, duong_dan: str, nguon_phien_am: str):
    """Gộp đoạn ngắn thành khối ~900 ký tự nhưng vẫn giữ mốc thời gian đầu/cuối."""
    tai_lieu = []
    dem: list[str] = []
    bat_dau = ket_thuc = 0.0
    do_dai = 0

    def chot():
        if not dem:
            return
        nhan = f"Phút {dinh_dang_thoi_gian(bat_dau)}–{dinh_dang_thoi_gian(ket_thuc)}"
        tai_lieu.append(Document(
            page_content=f"[{nhan}] " + " ".join(dem),
            metadata={
                "source": duong_dan,
                "context_label": nhan,
                "time_start": round(bat_dau, 2),
                "time_end": round(ket_thuc, 2),
                "chunk_mode": "theo_phan_tu",
                "nguon_phien_am": nguon_phien_am,
            },
        ))

    for doan in cac_doan:
        if not dem:
            bat_dau = doan.bat_dau
        dem.append(doan.noi_dung)
        ket_thuc = doan.ket_thuc
        do_dai += len(doan.noi_dung) + 1
        if do_dai >= DO_DAI_DOAN_PHIEN_AM:
            chot()
            dem, do_dai = [], 0
    chot()
    return tai_lieu


class MediaLoader:
    """
    Ưu tiên phụ đề .srt/.vtt đặt cạnh file media (chính xác hơn, tức thì);
    nếu không có thì phiên âm bằng faster-whisper và lưu cache.
    """

    def __init__(self, duong_dan: str, cho_phep_phien_am: bool = True):
        self.duong_dan = duong_dan
        self.cho_phep_phien_am = cho_phep_phien_am

    def load(self):
        goc = os.path.splitext(self.duong_dan)[0]
        for duoi in (".srt", ".vtt"):
            if os.path.exists(goc + duoi):
                return _gop_doan_phien_am(
                    doc_phu_de(goc + duoi), self.duong_dan, "phu_de_co_san"
                )

        cache = doc_cache(self.duong_dan)
        if cache is not None:
            return _gop_doan_phien_am(cache, self.duong_dan, "faster-whisper (cache)")
        if not self.cho_phep_phien_am:
            raise RuntimeError(
                "Video chưa có bản phiên âm. Chạy `python phien_am_video.py` để "
                "phiên âm trước, hoặc đặt RAG_TRANSCRIBE_ON_INDEX=1."
            )
        return _gop_doan_phien_am(
            phien_am(self.duong_dan), self.duong_dan, "faster-whisper"
        )


class PhuDeLoader:
    def __init__(self, duong_dan: str):
        self.duong_dan = duong_dan

    def load(self):
        return _gop_doan_phien_am(
            doc_phu_de(self.duong_dan), self.duong_dan, "phu_de_co_san"
        )


# ============================================================
# ẢNH RỜI (.jpg/.png) - ảnh chụp bảng, đề thi photo, trang sách chụp bằng điện thoại
# ============================================================
class AnhLoader:
    """
    Ảnh rời không có lớp văn bản nào để lui về như PDF: không OCR được thì tài
    liệu rỗng. Vì vậy thiếu Tesseract thì báo lỗi thẳng, thay vì trả Document
    trống rồi để file bị đánh dấu "không có chữ" mà người dùng không hiểu vì sao.
    """

    def __init__(self, duong_dan: str, cho_phep_ocr: bool = True):
        self.duong_dan = duong_dan
        self.cho_phep_ocr = cho_phep_ocr

    def load(self):
        import ocr_pdf

        if not self.cho_phep_ocr:
            return []
        van_ban = ocr_pdf.ocr_file_anh(self.duong_dan)
        if not van_ban.strip():
            return []
        return [Document(
            page_content=van_ban,
            metadata={
                "source": self.duong_dan,
                "nguon_van_ban": "ocr",
                "context_label": "Ảnh chụp",
            },
        )]


# ============================================================
# EPUB - sách điện tử: mỗi tệp XHTML bên trong là một phần
# ============================================================
def _van_ban_tu_html(noi_dung_html: str) -> str:
    """Gỡ thẻ nhưng giữ ranh giới khối, nếu không hai đoạn văn sẽ dính liền."""
    khong_script = re.sub(
        r"(?is)<(script|style|head)[^>]*>.*?</\1>", " ", noi_dung_html
    )
    xuong_dong = re.sub(
        r"(?i)<(br|/p|/div|/h[1-6]|/li|/tr|/td)[^>]*>", "\n", khong_script
    )
    van_ban = unescape(re.sub(r"(?s)<[^>]+>", " ", xuong_dong))
    van_ban = re.sub(r"[ \t\xa0]+", " ", van_ban)
    return re.sub(r"\n[ \t]*\n[ \t]*\n+", "\n\n", van_ban).strip()


def _tieu_de_chuong(noi_dung_html: str) -> str:
    for mau in (r"(?is)<h[1-3][^>]*>(.*?)</h[1-3]>", r"(?is)<title[^>]*>(.*?)</title>"):
        khop = re.search(mau, noi_dung_html)
        if khop:
            tieu_de = _van_ban_tu_html(khop.group(1))
            if tieu_de:
                return tieu_de[:80]
    return ""


class EpubLoader:
    """
    .epub là một gói zip chứa nhiều tệp XHTML. Đọc theo đúng thứ tự trong
    <spine> (thứ tự đọc thật của sách, không phải thứ tự tên file trong zip) và
    giữ mỗi tệp thành một Document riêng, nhờ vậy trích dẫn chỉ đúng chương thay
    vì cả quyển sách.
    """

    # Bỏ qua phần bìa/mục lục/bản quyền: gần như không có nội dung tra cứu được
    # nhưng lại rất khớp từ khóa tên sách nên hay chen lên đầu kết quả.
    DO_DAI_TOI_THIEU = 200

    def __init__(self, duong_dan: str):
        self.duong_dan = duong_dan

    @staticmethod
    def _ten_the(the: str) -> str:
        """'{http://www.idpf.org/2007/opf}item' -> 'item'."""
        return the.rsplit("}", 1)[-1]

    def _duong_dan_opf(self, goi, ElementTree) -> str:
        try:
            container = ElementTree.fromstring(goi.read("META-INF/container.xml"))
        except (KeyError, ElementTree.ParseError):
            container = None
        if container is not None:
            for el in container.iter():
                if self._ten_the(el.tag) == "rootfile" and el.get("full-path"):
                    return el.get("full-path")
        # Gói thiếu container.xml vẫn còn cứu được nếu tìm thấy file .opf.
        for ten in goi.namelist():
            if ten.lower().endswith(".opf"):
                return ten
        raise ValueError(
            "File .epub không có phần mô tả nội dung (.opf) nên không đọc được."
        )

    def load(self):
        import posixpath
        import zipfile
        from urllib.parse import unquote
        from xml.etree import ElementTree

        with zipfile.ZipFile(self.duong_dan) as goi:
            duong_dan_opf = self._duong_dan_opf(goi, ElementTree)
            opf = ElementTree.fromstring(goi.read(duong_dan_opf))
            thu_muc_goc = posixpath.dirname(duong_dan_opf)

            manifest = {}
            thu_tu_doc = []
            for el in opf.iter():
                ten = self._ten_the(el.tag)
                if ten == "item" and el.get("id") and el.get("href"):
                    manifest[el.get("id")] = el.get("href")
                elif ten == "itemref" and el.get("idref"):
                    thu_tu_doc.append(el.get("idref"))

            cac_tep = [manifest[ma] for ma in thu_tu_doc if ma in manifest]
            if not cac_tep:
                cac_tep = sorted(
                    href for href in manifest.values()
                    if href.lower().endswith((".xhtml", ".html", ".htm"))
                )

            ket_qua = []
            for thu_tu, href in enumerate(cac_tep, 1):
                trong_goi = posixpath.normpath(
                    posixpath.join(thu_muc_goc, unquote(href))
                )
                try:
                    noi_dung_goc = goi.read(trong_goi).decode("utf-8", errors="replace")
                except KeyError:
                    continue
                van_ban = _van_ban_tu_html(noi_dung_goc)
                if len(van_ban) < self.DO_DAI_TOI_THIEU:
                    continue
                tieu_de = _tieu_de_chuong(noi_dung_goc)
                ket_qua.append(Document(
                    page_content=van_ban,
                    metadata={
                        "source": self.duong_dan,
                        "chunk_mode": "theo_phan_tu",
                        "context_label": tieu_de or f"Phần {thu_tu}",
                    },
                ))

        if not ket_qua:
            raise ValueError("File .epub không có phần văn bản nào đọc được.")
        return ket_qua


# ============================================================
# ĐỊNH TUYẾN
# ============================================================
def tao_loader_cho_file(duong_dan: str, bao_tien_do_ocr=None):
    """
    .pdf   -> PyPDFLoader, tự chuyển sang OCR nếu là bản scan không có lớp chữ
    .docx  -> Docx2txtLoader     (UnstructuredFileLoader crash STATUS_STACK_BUFFER_OVERRUN)
    .doc   -> chuyển sang .docx bằng Word/LibreOffice rồi đọc
    .pptx  -> PptxLoader theo từng slide
    bảng   -> BangDuLieuLoader theo từng sheet, lặp header mỗi khối dòng
    .html  -> trafilatura
    .epub  -> EpubLoader theo từng phần trong <spine>
    ảnh    -> Tesseract OCR (ảnh chụp bảng, đề thi photo)
    media  -> phụ đề có sẵn hoặc phiên âm faster-whisper (có mốc thời gian)
    """
    duoi = os.path.splitext(duong_dan)[1].lower()
    if duoi == ".pdf":
        return PdfLoader(
            duong_dan,
            cho_phep_ocr=os.getenv("RAG_OCR_ON_INDEX", "1") == "1",
            bao_tien_do_ocr=bao_tien_do_ocr,
        )
    if duoi == ".docx":
        from langchain_community.document_loaders import Docx2txtLoader
        return Docx2txtLoader(duong_dan)
    if duoi == ".doc":
        return DocCuLoader(duong_dan)
    if duoi in DINH_DANG_TRINH_CHIEU:
        return PptxLoader(duong_dan)
    if duoi in DINH_DANG_BANG:
        return BangDuLieuLoader(duong_dan)
    if duoi in {".html", ".htm"}:
        return HtmlLoader(duong_dan)
    if duoi == ".epub":
        return EpubLoader(duong_dan)
    if duoi in DINH_DANG_ANH:
        return AnhLoader(
            duong_dan, cho_phep_ocr=os.getenv("RAG_OCR_ON_INDEX", "1") == "1"
        )
    if duoi in DINH_DANG_PHU_DE:
        return PhuDeLoader(duong_dan)
    if duoi in DINH_DANG_MEDIA:
        return MediaLoader(
            duong_dan,
            cho_phep_phien_am=os.getenv("RAG_TRANSCRIBE_ON_INDEX", "1") == "1",
        )
    return VanBanThuanLoader(duong_dan)
